# Copyright 2026 Seedian
# License AGPL-3.0 or later (http://www.gnu.org/licenses/agpl).

import fnmatch
import re
import zipfile
from datetime import date
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from lxml import etree
from openpyxl import load_workbook

from odoo.exceptions import AccessError, UserError
from odoo.fields import Command
from odoo.tests import TransactionCase

from odoo.addons.collection.models.collection_common import privileged

from .test_collection_acl import TestCollectionAcl
from .test_collection_payment import (
    create_internal_seller,
    dummy_customer_partner,
    ensure_enabled_seller,
    insert_company_like,
)

XLSX_MODEL = "collection.sheet.xlsx"
XLSX_MIMETYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_MERGES = {
    "A1:H1",
    "A2:H2",
    "A3:H3",
    "A4:H4",
    "A6:H6",
    "A7:B7",
    "C7:E7",
    "F7:H7",
    "A8:B8",
    "C8:E8",
    "F8:H8",
}
SHEET_FORM_XML = Path(__file__).resolve().parents[1] / "views" / "collection_views.xml"
QWEB_TEMPLATE = (
    Path(__file__).resolve().parents[1] / "report" / "collection_sheet_templates.xml"
)
XLSX_SOURCE = (
    Path(__file__).resolve().parents[1] / "report" / "collection_sheet_xlsx.py"
)
ACTION_SOURCE = (
    Path(__file__).resolve().parents[1] / "models" / "collection_sheet_report.py"
)


def _code_without_comments(path):
    source = path.read_text(encoding="utf-8")
    return "\n".join(
        line for line in source.splitlines() if not line.lstrip().startswith("#")
    )


def _is_same_date(value, expected):
    if value == expected:
        return True
    if getattr(value, "date", None) and callable(value.date):
        try:
            return value.date() == expected
        except TypeError:
            return False
    year = getattr(value, "year", None)
    month = getattr(value, "month", None)
    day = getattr(value, "day", None)
    return (year, month, day) == (expected.year, expected.month, expected.day)


class TestCollectionSheetXlsx(TransactionCase):
    """RED: Planilla XLSX zip, freeze, cached SUM, act_url, ACL."""

    def setUp(self):
        super().setUp()
        self._doc_seq = 0
        self.seller = create_internal_seller(self.env, "xlsx_seller_ana")
        self.seller.write({"name": "XLSX Seller Ana"})
        ensure_enabled_seller(self.env, self.seller)

    def _next_document_id(self, prefix="XLSX"):
        self._doc_seq += 1
        return f"{prefix}-{self._testMethodName}-{self._doc_seq}"

    def _create_sheet(self, *, seller=None, force_new=False, company=None, **extra):
        seller = seller or self.seller
        company = company or self.env.company
        ensure_enabled_seller(self.env, seller, company=company)
        extra.setdefault("company_id", company.id)
        if not force_new:
            existing = self.env["collection.sheet"].search(
                [
                    ("seller_id", "=", seller.id),
                    ("company_id", "=", extra["company_id"]),
                    ("state", "=", "open"),
                ],
                limit=1,
            )
            if existing:
                return existing
        extra.setdefault("seller_id", seller.id)
        return self.env["collection.sheet"].create(extra)

    def _create_voucher(self, partner, sheet, amount=100.0, **extra):
        ensure_enabled_seller(self.env, sheet.seller_id, company=sheet.company_id)
        vals = {
            "partner_id": partner.id,
            "document_type": extra.pop("document_type", "bank_transfer"),
            "document_id": extra.pop("document_id", self._next_document_id()),
            "seller_id": sheet.seller_id.id,
            "sheet_id": sheet.id,
            "company_id": sheet.company_id.id,
            "amount": amount,
        }
        vals.update(extra)
        if (
            vals.get("document_type")
            in (
                "check",
                "echeck",
            )
            and "check_number" not in extra
        ):
            # check_number is mandatory at save for check/echeck
            # (_check_check_number_required); a random default keeps
            # existing fixtures unique enough to dodge
            # _check_check_number_unique without spelling one out.
            vals["check_number"] = f"CHK-{uuid4().hex[:10]}"
        return self.env["collection.voucher"].create(vals)

    def _enable_seller_deposit_type(self, company=None):
        catalog = self.env.ref("collection.document_type_seller_deposit")
        (company or self.env.company).write(
            {"allowed_document_type_ids": [Command.link(catalog.id)]}
        )
        return catalog

    def _create_seller_deposit(self, sheet, amount=80.0, **extra):
        self._enable_seller_deposit_type(company=sheet.company_id)
        ensure_enabled_seller(self.env, sheet.seller_id, company=sheet.company_id)
        vals = {
            "document_type": "seller_deposit",
            "document_id": extra.pop("document_id", self._next_document_id("SD")),
            "seller_id": sheet.seller_id.id,
            "sheet_id": sheet.id,
            "company_id": sheet.company_id.id,
            "amount": amount,
            "voucher_number": extra.pop("voucher_number", "SD-1"),
            "voucher_date": extra.pop("voucher_date", date(2026, 9, 8)),
        }
        vals.update(extra)
        return self.env["collection.voucher"].create(vals)

    def _line_vals(self, amount_applied=100.0, invoice_balance=100.0, **extra):
        vals = {
            "external_invoice_id": extra.pop(
                "external_invoice_id", self._next_document_id("INV")
            ),
            "invoice_number": extra.pop("invoice_number", "A-1"),
            "issue_date": extra.pop("issue_date", date(2026, 1, 1)),
            "due_date": extra.pop("due_date", date(2026, 1, 31)),
            "invoice_amount": extra.pop("invoice_amount", invoice_balance),
            "invoice_balance": invoice_balance,
            "invoice_state": extra.pop("invoice_state", "open"),
            "amount_applied": amount_applied,
            "document_kind": extra.pop("document_kind", "invoice"),
        }
        vals.update(extra)
        return vals

    def _create_allocation(self, sheet, partner, vouchers=None, lines=None, **extra):
        vals = {
            "sheet_id": sheet.id,
            "partner_id": partner.id,
        }
        if vouchers:
            vals["voucher_ids"] = [Command.set(vouchers.ids)]
        if lines is not None:
            vals["line_ids"] = [Command.create(line) for line in lines]
        vals.update(extra)
        return self.env["collection.allocation"].create(vals)

    def _populated_sheet(
        self,
        *,
        seller=None,
        company=None,
        partner_name="Cliente Norte SA",
        invoice_number="FAC-1001",
        amount=100.0,
        document_type="bank_transfer",
        commission_percent=0.0,
        commission_withheld=0.0,
    ):
        seller = seller or self.seller
        partner = dummy_customer_partner(self.env, customer=partner_name)
        sheet = self._create_sheet(
            seller=seller,
            force_new=True,
            company=company,
            commission_percent=commission_percent,
            commission_withheld=commission_withheld,
        )
        voucher = self._create_voucher(
            partner, sheet, amount=amount, document_type=document_type
        )
        self._create_allocation(
            sheet,
            partner,
            vouchers=voucher,
            lines=[
                self._line_vals(
                    amount_applied=amount,
                    invoice_balance=amount,
                    invoice_number=invoice_number,
                )
            ],
        )
        return sheet, partner

    def _render_xlsx(self, sheet):
        self.assertIn(XLSX_MODEL, self.env)
        payload = sheet._collection_report_data()
        content = self.env[XLSX_MODEL]._render_xlsx(payload)
        self.assertIsInstance(content, (bytes, bytearray), type(content))
        return bytes(content), payload

    def _load(self, content, data_only=False):
        return load_workbook(BytesIO(content), data_only=data_only)

    def _sheet_form_root(self):
        tree = etree.parse(str(SHEET_FORM_XML))
        records = tree.xpath("//record[@id='view_collection_sheet_form']")
        self.assertTrue(records, "view_collection_sheet_form must exist")
        forms = records[0].xpath(".//form")
        self.assertTrue(forms, "sheet form arch must contain <form>")
        return forms[0]

    def _create_partner(self, name):
        Partner = self.env["res.partner"].with_context(
            tracking_disable=True,
            mail_create_nolog=True,
            mail_notrack=True,
            mail_create_nosubscribe=True,
        )
        vals = {"name": name}
        if "autopost_bills" in Partner._fields:
            vals["autopost_bills"] = "ask"
            return Partner.create(vals)
        self.env.cr.execute(
            """
            SELECT 1
              FROM information_schema.columns
             WHERE table_name = 'res_partner'
               AND column_name = 'autopost_bills'
            """
        )
        if not self.env.cr.fetchone():
            return Partner.create(vals)
        self.env.cr.execute(
            """
            INSERT INTO res_partner (
                name, email, active, complete_name, type, autopost_bills,
                create_uid, write_uid, create_date, write_date
            )
            VALUES (
                %s, %s, true, %s, 'contact', 'ask',
                1, 1,
                (now() AT TIME ZONE 'UTC'),
                (now() AT TIME ZONE 'UTC')
            )
            RETURNING id
            """,
            (name, f"{name}@example.com", name),
        )
        return Partner.browse(self.env.cr.fetchone()[0])

    def _create_company(self, name):
        partner = self._create_partner(name)
        company = insert_company_like(self.env, name, partner.id)
        company._collection_copy_sequences()
        company._collection_enable_all_document_types()
        return company

    def _create_user(self, login, group_xmlid):
        return TestCollectionAcl._create_user.__func__(self, login, group_xmlid)

    def _partner_has_autopost_bills_column(self):
        return TestCollectionAcl._partner_has_autopost_bills_column.__func__(self)

    def _restrict_user_to_company(self, user, company):
        user.write({"company_ids": [Command.link(company.id)]})
        user.write(
            {
                "company_id": company.id,
                "company_ids": [Command.set([company.id])],
            }
        )

    def test_xlsx_bytes_are_zip_with_workbook(self):
        """XLSX bytes MUST start PK and the zip MUST contain xl/workbook.xml."""
        sheet, _partner = self._populated_sheet()
        content, _payload = self._render_xlsx(sheet)
        self.assertTrue(content.startswith(b"PK"), content[:8])
        with zipfile.ZipFile(BytesIO(content)) as zipped:
            names = zipped.namelist()
        self.assertIn("xl/workbook.xml", names)

    def test_xlsx_print_setup_freeze_and_column_width(self):
        """Worksheet MUST freeze A6, landscape paper 9, fit width, titles 1:4."""
        sheet, _partner = self._populated_sheet()
        content, _payload = self._render_xlsx(sheet)
        workbook = self._load(content)
        worksheet = workbook.active
        self.assertEqual(worksheet.title, "Planilla de Cobranza")
        self.assertEqual(worksheet.freeze_panes, "A6")
        self.assertEqual(worksheet.page_setup.orientation, "landscape")
        self.assertEqual(worksheet.page_setup.paperSize, 9)
        self.assertEqual(worksheet.page_setup.fitToWidth, 1)
        title_rows = (worksheet.print_title_rows or "").replace("$", "")
        self.assertEqual(title_rows, "1:4")
        self.assertEqual(worksheet.column_dimensions["B"].width, 32)

    def test_xlsx_header_has_title_but_no_company_nor_printed_date(self):
        """Header MUST show title, reference, seller and period; no company/Impreso."""
        sheet, _partner = self._populated_sheet()
        sheet.write({"date_from": date(2026, 10, 1), "date_to": date(2026, 10, 15)})
        content, payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        self.assertEqual(worksheet["A1"].value, "Planilla de Cobranza")
        self.assertEqual(worksheet["A2"].value, payload["reference"])
        self.assertEqual(
            worksheet["A4"].value,
            f"Vendedor: {payload['seller_name']} · Período: 01/10/2026 – 15/10/2026",
        )
        company_name = sheet.company_id.name
        values = [
            str(cell.value)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        self.assertNotIn("Impreso el", values)
        self.assertNotIn(company_name, values)
        self.assertFalse(
            any("Impreso" in value for value in values),
            "printed date label must not render",
        )
        self.assertIn("Vendedor", values)

    def test_xlsx_merges_and_number_formats(self):
        """Header merges and money/percent formats MUST match the template."""
        sheet, _partner = self._populated_sheet(commission_percent=4.5)
        content, _payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        merges = {str(item) for item in worksheet.merged_cells.ranges}
        missing = HEADER_MERGES - merges
        self.assertFalse(missing, f"missing merges: {sorted(missing)}")
        formats = {
            cell.number_format
            for row in worksheet.iter_rows()
            for cell in row
            if cell.number_format and cell.number_format != "General"
        }
        self.assertTrue(
            any("#,##0.00" in (fmt or "") for fmt in formats),
            formats,
        )
        self.assertTrue(
            any("0.0%" in (fmt or "") for fmt in formats),
            formats,
        )
        self.assertTrue(
            any("$" in (fmt or "") for fmt in formats),
            formats,
        )

    def test_xlsx_cached_subtotals_match_payload(self):
        """data_only cached SUM cells MUST equal payload totals and not be None."""
        sheet, _partner = self._populated_sheet(amount=125.5, commission_percent=4.5)
        content, payload = self._render_xlsx(sheet)
        formulas = self._load(content).active
        cached = self._load(content, data_only=True).active
        formula_coords = []
        for row in formulas.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and "SUM" in value.upper():
                    formula_coords.append(cell.coordinate)
                    cached_value = cached[cell.coordinate].value
                    self.assertIsNotNone(
                        cached_value,
                        f"{cell.coordinate} formula {value} cached None",
                    )
        self.assertTrue(formula_coords, "workbook MUST contain SUM formulas")
        cached_numbers = []
        for row in cached.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cached_numbers.append(float(cell.value))

        def _assert_cached(amount, label):
            matches = [
                number for number in cached_numbers if abs(number - amount) < 0.0001
            ]
            self.assertTrue(matches, f"{label}={amount} missing from cached cells")

        _assert_cached(payload["collected_total"], "collected_total")
        _assert_cached(payload["applied_total"], "applied_total")
        _assert_cached(payload["customer_totals"]["debt"], "customer debt")
        _assert_cached(payload["customer_totals"]["collected"], "customer collected")
        _assert_cached(payload["recon"]["cash_due"], "recon cash due")

    def test_xlsx_percent_scale_4_5_is_excel_fraction(self):
        """Odoo commission_percent 4.5 MUST be stored as 0.045 with 0.0% format."""
        sheet, _partner = self._populated_sheet(commission_percent=4.5)
        content, payload = self._render_xlsx(sheet)
        self.assertEqual(payload["commission_percent"], 4.5)
        worksheet = self._load(content).active
        percent_cells = []
        for row in worksheet.iter_rows():
            for cell in row:
                fmt = cell.number_format or ""
                if "0.0%" in fmt and cell.value is not None:
                    percent_cells.append(cell)
        self.assertTrue(percent_cells, "workbook MUST contain a 0.0% cell")
        for cell in percent_cells:
            self.assertAlmostEqual(cell.value, 0.045, places=6, msg=cell.coordinate)

    def test_xlsx_recon_reads_as_the_cash_calculation(self):
        """Recon MUST read as cash collected minus deposits and commission kept."""
        sheet, _partner = self._populated_sheet(
            amount=200.0,
            document_type="cash",
            commission_percent=10.0,
            commission_withheld=5.0,
        )
        content, payload = self._render_xlsx(sheet)
        recon = payload["recon"]
        self.assertEqual(recon["cash_due"], 195.0)
        worksheet = self._load(content).active
        cached = self._load(content, data_only=True).active
        rows = []
        for label, amount in (
            ("Efectivo cobrado", 200.0),
            ("(−) Depositado por el Vendedor", 0.0),
            ("(−) Comisión Retenida por el Vendedor", 5.0),
            ("(=) Efectivo a entregar", 195.0),
            ("Cobrado por otros medios", 0.0),
            ("Comisión Devengada", 20.0),
            ("(=) Comisión a Pagar al Vendedor", 15.0),
        ):
            with self.subTest(label=label):
                # The verdict repeats some labels in other columns; recon is column A.
                cell = next(
                    (c for (c,) in worksheet.iter_rows(max_col=1) if c.value == label),
                    None,
                )
                self.assertIsNotNone(cell, label)
                rows.append(cell.row)
                self.assertAlmostEqual(
                    float(cached.cell(cell.row, 7).value), amount, places=2
                )
        self.assertEqual(rows, sorted(rows))
        due = self._cell_with_value(worksheet, "(=) Efectivo a entregar")
        self.assertTrue(str(worksheet.cell(due.row, 7).value).startswith("="))

    def test_xlsx_recon_follows_the_pdf_order(self):
        """Cash calculation, then composition, then commission, as in the PDF."""
        sheet, _partner = self._populated_sheet(document_type="cash")
        content, payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        rows = []
        for label in (
            "Conciliación del Vendedor",
            "(=) Efectivo a entregar",
            "Cobrado por otros medios",
            "Composición de lo Entregado",
            "Comisión del Vendedor",
            "(=) Comisión a Pagar al Vendedor",
        ):
            cell = next(
                (c for (c,) in worksheet.iter_rows(max_col=1) if c.value == label),
                None,
            )
            self.assertIsNotNone(cell, label)
            rows.append(cell.row)
        self.assertEqual(rows, sorted(rows))
        # The PDF composition has no total; neither does the XLSX.
        self.assertIsNone(self._cell_with_value(worksheet, "Entregado por el Vendedor"))
        header = self._cell_with_value(worksheet, "Depósitos del vendedor")
        self.assertIsNone(worksheet.cell(header.row, header.column + 1).value)
        self.assertEqual(payload["tendered"]["cash"], payload["recon"]["cash_due"])

    def test_xlsx_has_no_banner_or_difference(self):
        """The always-zero Diferencia and its status banner MUST be gone."""
        sheet, _partner = self._populated_sheet()
        content, _payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        for text in ("(=) Diferencia", "Diferencia", "Cuadrada"):
            with self.subTest(text=text):
                self.assertIsNone(self._cell_with_value(worksheet, text))
        self.assertNotIn("A9:H9", {str(r) for r in worksheet.merged_cells.ranges})

    def test_xlsx_tendered_shows_echecks_retentions_seller_deposits(self):
        """XLSX recon header MUST show the new tendered buckets and amounts."""
        seller = create_internal_seller(self.env, "xlsx_tendered_buckets")
        ensure_enabled_seller(self.env, seller)
        sheet, partner = self._populated_sheet(
            seller=seller,
            partner_name="Tendered Buckets SA",
            document_type="bank_transfer",
            amount=100.0,
        )
        echeck = self._create_voucher(
            partner, sheet, amount=30.0, document_type="echeck"
        )
        retention = self._create_voucher(
            partner, sheet, amount=12.0, document_type="retention_certificate"
        )
        remittance = self._create_seller_deposit(sheet, amount=80.0)
        content, payload = self._render_xlsx(sheet)
        tendered = payload["tendered"]
        self.assertEqual(tendered["echecks"], echeck.amount)
        self.assertEqual(tendered["retentions"], retention.amount)
        self.assertEqual(tendered["seller_deposits"], remittance.amount)
        worksheet = self._load(content).active
        cached = self._load(content, data_only=True).active
        for header, amount in (
            ("Echeqs", echeck.amount),
            ("Retenciones", retention.amount),
            ("Depósitos del vendedor", remittance.amount),
            ("Efectivo en Caja", tendered["cash"]),
            ("Transferencias", tendered["transfers"]),
        ):
            title = self._cell_with_value(worksheet, header)
            self.assertIsNotNone(title, f"missing XLSX header {header!r}")
            below = cached.cell(title.row + 1, title.column)
            self.assertAlmostEqual(float(below.value), amount, places=2, msg=header)

    def test_qweb_recon_shows_new_tendered_buckets(self):
        """QWeb recon table MUST t-out echecks, retentions, and seller_deposits."""
        self.assertTrue(QWEB_TEMPLATE.is_file(), QWEB_TEMPLATE)
        arch = QWEB_TEMPLATE.read_text(encoding="utf-8")
        for key in ("echecks", "retentions", "seller_deposits"):
            self.assertIn(f"payload['tendered']['{key}']", arch, key)
        self.assertIn("Echeqs", arch)
        self.assertIn("Retenciones", arch)
        self.assertIn("Depósitos del vendedor", arch)

    def test_qweb_recon_shows_the_commission_owed_to_the_seller(self):
        """QWeb recon MUST print the commission owed to the seller."""
        arch = QWEB_TEMPLATE.read_text(encoding="utf-8")
        self.assertIn("payload['recon']['commission_owed']", arch)
        self.assertIn("Comisión a Pagar al Vendedor", arch)

    def test_xlsx_recon_shows_the_commission_owed_to_the_seller(self):
        """XLSX recon MUST show the commission minus what was kept from cash."""
        sheet, _partner = self._populated_sheet(
            amount=200.0, commission_percent=10.0, commission_withheld=5.0
        )
        content, payload = self._render_xlsx(sheet)
        self.assertEqual(payload["recon"]["commission_owed"], 15.0)
        worksheet = self._load(content).active
        cached = self._load(content, data_only=True).active
        label = self._cell_with_value(worksheet, "(=) Comisión a Pagar al Vendedor")
        self.assertIsNotNone(label)
        self.assertAlmostEqual(
            float(cached.cell(label.row, label.column + 6).value), 15.0, places=2
        )

    def test_action_print_xlsx_act_url_and_attachment(self):
        """action_print_xlsx MUST return act_url download and one xlsx attachment."""
        sheet, _partner = self._populated_sheet()
        action = sheet.action_print_xlsx()
        self.assertEqual(action["type"], "ir.actions.act_url")
        self.assertEqual(action["target"], "download")
        self.assertTrue(action["url"].startswith("/web/content/"))
        self.assertTrue(action["url"].endswith("?download=true"))
        attachment_id = int(action["url"].split("/")[3].split("?")[0])
        attachment = self.env["ir.attachment"].browse(attachment_id)
        self.assertTrue(attachment.exists())
        self.assertEqual(attachment.res_model, "collection.sheet")
        self.assertEqual(attachment.res_id, sheet.id)
        self.assertEqual(attachment.mimetype, XLSX_MIMETYPE)
        self.assertTrue(attachment.name.endswith(".xlsx"))
        self.assertNotIn("/", attachment.name)
        raw = attachment.raw or b""
        self.assertTrue(raw.startswith(b"PK"), raw[:8])

    def test_xlsx_in_all_five_states(self):
        """Collection Sheet XLSX MUST download in every state, with no UserError."""
        sheet, _partner = self._populated_sheet()
        if sheet.voucher_ids:
            with privileged("collection_audited_state"):
                sheet.voucher_ids.write({"state": "validated_collections"})

        def _seller_close_and_declare():
            sheet.action_seller_close()
            # Zero withheld: Confirm mint is a no-op so this walk stays a
            # state/download check and does not require replaceable cash.
            sheet.action_declare_cash(commission_withheld=0)

        advances = (
            ("open", lambda: None),
            ("submitted", _seller_close_and_declare),
            ("proposal", sheet.action_confirm),
            ("closed", sheet.action_close),
            ("cancelled", sheet.action_cancel),
        )
        for target, advance in advances:
            with self.subTest(state=target):
                advance()
                self.assertEqual(sheet.state, target)
                action = sheet.action_print_xlsx()
                self.assertEqual(action["type"], "ir.actions.act_url")
                self.assertTrue(action["url"].endswith("?download=true"))
                attachment_id = int(action["url"].split("/")[3].split("?")[0])
                raw = self.env["ir.attachment"].browse(attachment_id).raw or b""
                self.assertTrue(raw.startswith(b"PK"), raw[:8])

    def test_unsaved_xlsx_refused(self):
        """Unsaved sheets MUST raise UserError and MUST NOT emit an XLSX download."""
        sheet = self.env["collection.sheet"].new(
            {"seller_id": self.seller.id, "company_id": self.env.company.id}
        )
        with self.assertRaises(UserError) as error:
            sheet.action_print_xlsx()
        self.assertIn("saved", str(error.exception).lower())

    def test_action_print_xlsx_second_call_replaces_attachment(self):
        """A second Export MUST replace the same-name attachment, not duplicate."""
        sheet, _partner = self._populated_sheet()
        first = sheet.action_print_xlsx()
        second = sheet.action_print_xlsx()
        filename = f"{(sheet.name or '').replace('/', '-')}.xlsx"
        attachments = self.env["ir.attachment"].search(
            [
                ("res_model", "=", "collection.sheet"),
                ("res_id", "=", sheet.id),
                ("name", "=", filename),
            ]
        )
        self.assertEqual(len(attachments), 1)
        self.assertEqual(attachments.mimetype, XLSX_MIMETYPE)
        first_id = int(first["url"].split("/")[3].split("?")[0])
        second_id = int(second["url"].split("/")[3].split("?")[0])
        self.assertEqual(attachments.id, second_id)
        self.assertFalse(self.env["ir.attachment"].browse(first_id).exists())

    def test_header_xlsx_button_last_before_state(self):
        """Collection Sheet XLSX MUST be the last header object button before state."""
        form = self._sheet_form_root()
        header = form.xpath("./header")
        self.assertTrue(header, "sheet form must have a header")
        buttons = [node for node in header[0].xpath("./*") if node.tag == "button"]
        names = [node.get("name") for node in buttons]
        self.assertIn("action_open_document_review", names)
        self.assertNotIn("action_open_voucher_reassign", names)
        self.assertNotIn("action_reopen", names)
        self.assertEqual(names[-1], "action_print_xlsx")
        xlsx = buttons[-1]
        self.assertEqual(xlsx.get("type"), "object")
        self.assertEqual(xlsx.get("string"), "Collection Sheet XLSX")
        following = xlsx.getnext()
        self.assertIsNotNone(following)
        self.assertEqual(following.tag, "field")
        self.assertEqual(following.get("name"), "state")

    def test_sheet_form_three_notebook_pages_recon_second_group(self):
        """Notebook MUST have three pages; recon fields stay in a second inner group."""
        form = self._sheet_form_root()
        notebooks = form.xpath(".//notebook")
        self.assertEqual(len(notebooks), 1)
        pages = notebooks[0].xpath("./page")
        self.assertEqual(
            [page.get("string") for page in pages],
            ["Allocations", "Vouchers", "Debt"],
        )
        groups = form.xpath("./sheet/group")
        self.assertTrue(groups, "sheet body must have an outer group")
        inner = groups[0].xpath("./group")
        self.assertGreaterEqual(
            len(inner),
            2,
            "recon fields MUST live in a second inner group, not a new page",
        )
        recon_names = [node.get("name") for node in inner[1].xpath("./field")]
        for field_name in (
            "commission_percent",
            "commission_amount",
            "cash_on_hand",
            "commission_withheld",
            "commission_owed",
        ):
            self.assertIn(field_name, recon_names, field_name)

    def test_company_b_access_error_on_xlsx(self):
        """Company-B users MUST get AccessError on action_print_xlsx with no file."""
        user_a = self._create_user("xlsx_company_a", "collection.group_collection_user")
        self._restrict_user_to_company(user_a, self.env.company)

        company_b = self._create_company("XLSX Hidden Co B")
        seller_b = create_internal_seller(self.env, "xlsx_hidden_seller")
        seller_b.write({"name": "Hidden Seller B"})
        sheet_b, _partner = self._populated_sheet(
            seller=seller_b,
            company=company_b,
            partner_name="Hidden Cliente XLSX",
        )
        hidden = self.env["collection.sheet"].with_user(user_a).browse(sheet_b.id)
        before = self.env["ir.attachment"].search(
            [
                ("res_model", "=", "collection.sheet"),
                ("res_id", "=", sheet_b.id),
            ]
        )
        with self.assertRaises(AccessError):
            hidden.action_print_xlsx()
        after = self.env["ir.attachment"].search(
            [
                ("res_model", "=", "collection.sheet"),
                ("res_id", "=", sheet_b.id),
            ]
        )
        self.assertEqual(after, before)

    def test_xlsx_sources_have_no_sudo(self):
        """XLSX writer and action_print_xlsx MUST NOT call sudo()."""
        self.assertTrue(XLSX_SOURCE.is_file(), XLSX_SOURCE)
        xlsx_code = _code_without_comments(XLSX_SOURCE)
        self.assertNotIn(".sudo(", xlsx_code)
        self.assertNotIn("sudo()", xlsx_code)
        action_code = _code_without_comments(ACTION_SOURCE)
        self.assertNotIn(".sudo(", action_code)
        self.assertNotIn("sudo()", action_code)

    def _cell_with_value(self, worksheet, expected):
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value == expected:
                    return cell
        return None

    def _cell_with_date(self, worksheet, expected):
        for row in worksheet.iter_rows():
            for cell in row:
                if expected and _is_same_date(cell.value, expected):
                    return cell
        return None

    def test_xlsx_header_currency_and_invoice_balance(self):
        """XLSX MUST print currency_name next to Moneda and invoice saldo."""
        partner = dummy_customer_partner(self.env, customer="XLSX-SALDO")
        sheet = self._create_sheet(force_new=True)
        voucher = self._create_voucher(partner, sheet, amount=80.0)
        self._create_allocation(
            sheet,
            partner,
            vouchers=voucher,
            lines=[
                self._line_vals(
                    amount_applied=80.0,
                    invoice_balance=321.5,
                    invoice_number="FAC-SALDO",
                    due_date=date(2026, 6, 15),
                )
            ],
        )
        content, payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        moneda = self._cell_with_value(worksheet, "Moneda")
        self.assertIsNotNone(moneda)
        currency_cell = self._cell_with_value(worksheet, payload["currency_name"])
        self.assertIsNotNone(currency_cell)
        self.assertEqual(currency_cell.row, moneda.row)
        saldo_header = self._cell_with_value(worksheet, "Saldo")
        self.assertIsNotNone(saldo_header)
        saldo = worksheet.cell(saldo_header.row + 1, saldo_header.column)
        self.assertAlmostEqual(float(saldo.value), 321.5, places=2)
        debt_header = self._cell_with_value(worksheet, "Saldo Total")
        self.assertIsNotNone(debt_header)
        cobrado_header = self._cell_with_value(worksheet, "Cobrado")
        self.assertIsNotNone(cobrado_header)
        self.assertEqual(debt_header.row, cobrado_header.row)
        self.assertEqual(debt_header.column, cobrado_header.column - 1)
        debt = worksheet.cell(debt_header.row + 1, debt_header.column)
        cobrado = worksheet.cell(debt_header.row + 1, debt_header.column + 1)
        aplicado = worksheet.cell(debt_header.row + 1, debt_header.column + 2)
        remanente = worksheet.cell(debt_header.row + 1, debt_header.column + 3)
        self.assertIn("SUMIF", str(debt.value).upper())
        self.assertIn("SUMIF", str(cobrado.value).upper())
        self.assertIn("SUMIF", str(aplicado.value).upper())
        self.assertRegex(str(remanente.value).replace("$", ""), r"^=C\d+-E\d+$")
        invoice_number = self._cell_with_value(worksheet, "FAC-SALDO")
        self.assertIsNotNone(invoice_number)
        self.assertIn(str(invoice_number.row), str(debt.value))
        self.assertIn("$E$", str(debt.value).upper())
        self.assertIn("$H$", str(cobrado.value).upper())
        self.assertIn("$F$", str(aplicado.value).upper())
        cached = self._load(content, data_only=True).active
        self.assertAlmostEqual(
            float(cached.cell(debt.row, debt.column).value), 321.5, places=2
        )
        self.assertAlmostEqual(
            float(cached.cell(cobrado.row, cobrado.column).value), 80.0, places=2
        )
        self.assertAlmostEqual(
            float(cached.cell(aplicado.row, aplicado.column).value), 80.0, places=2
        )
        self.assertAlmostEqual(
            float(cached.cell(remanente.row, remanente.column).value), 241.5, places=2
        )
        resumen_rem = self._cell_with_value(worksheet, "Deuda sin cobrar")
        self.assertIsNotNone(resumen_rem)
        resumen_formula = str(
            worksheet.cell(resumen_rem.row, resumen_rem.column + 2).value or ""
        )
        self.assertNotIn("C11-G11", resumen_formula.replace("$", ""))
        self.assertIn("=", resumen_formula)

    def test_xlsx_transfer_and_check_comprobante_fields(self):
        """XLSX MUST print voucher_number / check details, never the COL name."""
        partner = dummy_customer_partner(self.env, customer="XLSX-COMP")
        sheet = self._create_sheet(force_new=True)
        transfer = self._create_voucher(
            partner,
            sheet,
            amount=10.0,
            document_type="bank_transfer",
            voucher_number="TRF-XLSX",
        )
        check = self._create_voucher(
            partner,
            sheet,
            amount=20.0,
            document_type="check",
            check_number="CHK-XLSX",
            check_bank="Banco Galicia",
            check_type="common",
            check_issue_date=date(2026, 3, 1),
            check_payment_date=date(2026, 4, 15),
        )
        self._create_allocation(
            sheet,
            partner,
            vouchers=transfer | check,
            lines=[self._line_vals(amount_applied=30.0, invoice_balance=30.0)],
        )
        content, _payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        self.assertIsNotNone(self._cell_with_value(worksheet, "TRF-XLSX"))
        self.assertIsNone(self._cell_with_value(worksheet, transfer.name))
        self.assertIsNotNone(self._cell_with_value(worksheet, "CHK-XLSX"))
        self.assertIsNone(self._cell_with_value(worksheet, check.name))
        self.assertIsNotNone(self._cell_with_value(worksheet, "Banco Galicia"))
        self.assertIsNotNone(self._cell_with_value(worksheet, "Común"))
        self.assertIsNotNone(self._cell_with_date(worksheet, date(2026, 4, 15)))

    def test_xlsx_echeck_block_prints_check_columns(self):
        """XLSX echeck block MUST use the check columns, not the merged number."""
        partner = dummy_customer_partner(self.env, customer="XLSX-ECHK")
        sheet = self._create_sheet(force_new=True)
        echeck = self._create_voucher(
            partner,
            sheet,
            amount=20.0,
            document_type="echeck",
            check_number="ECHK-XLSX",
            check_bank="Banco Macro",
            check_type="deferred",
            check_issue_date=date(2026, 3, 2),
            check_payment_date=date(2026, 5, 20),
        )
        self._create_allocation(
            sheet,
            partner,
            vouchers=echeck,
            lines=[self._line_vals(amount_applied=20.0, invoice_balance=20.0)],
        )
        content, _payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        number = self._cell_with_value(worksheet, "ECHK-XLSX")
        self.assertIsNotNone(number)
        self.assertIsNone(self._cell_with_value(worksheet, echeck.name))
        self.assertFalse(
            any(number.coordinate in merged for merged in worksheet.merged_cells.ranges)
        )
        self.assertIsNotNone(self._cell_with_value(worksheet, "Banco Macro"))
        self.assertIsNotNone(self._cell_with_value(worksheet, "Diferido"))
        self.assertIsNotNone(self._cell_with_date(worksheet, date(2026, 5, 20)))
        self.assertIsNotNone(self._cell_with_value(worksheet, "Nº Cheque"))

    def test_xlsx_header_values_right_and_table_dates_centered(self):
        """Header field values MUST be right-aligned; table dates centered."""
        partner = dummy_customer_partner(self.env, customer="XLSX-ALIGN")
        sheet = self._create_sheet(force_new=True)
        voucher = self._create_voucher(
            partner,
            sheet,
            amount=50.0,
            document_type="bank_transfer",
            voucher_number="TRF-ALIGN",
        )
        self._create_allocation(
            sheet,
            partner,
            vouchers=voucher,
            lines=[
                self._line_vals(
                    amount_applied=50.0,
                    invoice_balance=50.0,
                    due_date=date(2026, 7, 20),
                )
            ],
        )
        content, payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        moneda_value = self._cell_with_value(worksheet, payload["currency_name"])
        self.assertIsNotNone(moneda_value)
        self.assertEqual(moneda_value.alignment.horizontal, "right")
        seller = self._cell_with_value(worksheet, payload["seller_name"])
        self.assertIsNotNone(seller)
        self.assertEqual(seller.alignment.horizontal, "right")
        due = self._cell_with_date(worksheet, date(2026, 7, 20))
        self.assertIsNotNone(due)
        self.assertEqual(due.alignment.horizontal, "center")
        voucher_date = payload["voucher_blocks"][0]["rows"][0]["date"]
        date_cell = self._cell_with_date(worksheet, voucher_date)
        self.assertIsNotNone(date_cell)
        self.assertEqual(date_cell.alignment.horizontal, "center")

    def test_xlsx_free_text_is_written_as_text_never_as_formula(self):
        """Names, numbers and banks starting with ``=`` stay inert strings."""
        injected = '=HYPERLINK("http://evil.example/x","click")'
        self.seller.write({"name": "=1+1"})
        partner = self._create_partner(injected)
        sheet = self._create_sheet(force_new=True)
        transfer = self._create_voucher(
            partner, sheet, amount=10.0, voucher_number="=2+2"
        )
        check = self._create_voucher(
            partner,
            sheet,
            amount=20.0,
            document_type="check",
            check_number="=3+3",
            check_bank="http://evil.example/bank",
        )
        self._create_allocation(
            sheet,
            partner,
            vouchers=transfer | check,
            lines=[
                self._line_vals(
                    amount_applied=30.0, invoice_balance=30.0, invoice_number="=4+4"
                )
            ],
        )
        content, _payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        for text in (injected, "=1+1", "=2+2", "=3+3", "=4+4"):
            with self.subTest(text=text):
                cells = [
                    cell
                    for row in worksheet.iter_rows()
                    for cell in row
                    if cell.value == text
                ]
                self.assertTrue(cells, f"{text!r} must be written")
                for cell in cells:
                    self.assertEqual(cell.data_type, "s", cell.coordinate)
        bank = self._cell_with_value(worksheet, "http://evil.example/bank")
        self.assertIsNotNone(bank)
        self.assertIsNone(bank.hyperlink)
        formulas = [
            str(cell.value)
            for row in worksheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        ]
        self.assertTrue(formulas, "intentional formulas must still be formulas")
        self.assertFalse([value for value in formulas if "HYPERLINK" in value])

    def _evaluate_sumifs(self, worksheet, formula):
        """Evaluate ``=SUMIF(range,criteria,sum_range)+...`` the way Excel does.

        Text criteria are case-insensitive and honour the ``*`` / ``?`` wildcards;
        numbers compare by value. Enough to recalculate the customer totals.
        """
        cell = r"\$?([A-Z]+)\$?(\d+)"
        part = re.compile(
            rf"SUMIF\({cell}:{cell},{cell},{cell}:{cell}\)", re.IGNORECASE
        )
        total = 0.0
        for match in part.finditer(formula):
            key_col, first, _key_col2, last = match.group(1, 2, 3, 4)
            crit_col, crit_row = match.group(5, 6)
            sum_col = match.group(7)
            criteria = worksheet[f"{crit_col}{crit_row}"].value
            for row in range(int(first), int(last) + 1):
                value = worksheet[f"{key_col}{row}"].value
                if isinstance(criteria, str):
                    pattern = criteria.lower().replace("[", "[[]")
                    hit = isinstance(value, str) and fnmatch.fnmatchcase(
                        value.lower(), pattern
                    )
                else:
                    hit = value is not None and value == criteria
                if hit:
                    total += float(worksheet[f"{sum_col}{row}"].value or 0.0)
        return total

    def test_xlsx_customer_totals_never_mix_customers(self):
        """Same names, or a ``*`` in a name, must not pull another customer's rows."""
        sheet = self._create_sheet(force_new=True)
        figures = (
            ("Cliente Duplicado", 100.0, 300.0),
            ("Cliente Duplicado", 40.0, 70.0),
            ("Cliente*", 25.0, 90.0),
        )
        for index, (name, collected, balance) in enumerate(figures):
            partner = self._create_partner(name)
            voucher = self._create_voucher(partner, sheet, amount=collected)
            self._create_allocation(
                sheet,
                partner,
                vouchers=voucher,
                lines=[
                    self._line_vals(
                        amount_applied=collected,
                        invoice_balance=balance,
                        invoice_number=f"FAC-MIX-{index}",
                    )
                ],
            )
        content, payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        self.assertTrue(worksheet.column_dimensions["L"].hidden, "key column")
        header = self._cell_with_value(worksheet, "Saldo Total")
        self.assertIsNotNone(header)
        customers = payload["customers"]
        self.assertEqual(len(customers), 3)
        for offset, expected in enumerate(customers):
            row = header.row + 1 + offset
            with self.subTest(customer=expected["name"], row=row):
                for col, key in ((3, "debt"), (4, "collected"), (5, "applied")):
                    formula = str(worksheet.cell(row, col).value)
                    self.assertTrue(formula.startswith("="), formula)
                    self.assertAlmostEqual(
                        self._evaluate_sumifs(worksheet, formula),
                        expected[key],
                        places=2,
                        msg=f"{key}: {formula}",
                    )

    def test_xlsx_verdict_comes_first(self):
        """Efectivo a entregar, otros medios and the commission owed MUST lead."""
        sheet, _partner = self._populated_sheet(amount=125.5)
        content, payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        cached = self._load(content, data_only=True).active
        datos = self._cell_with_value(worksheet, "Datos de la Planilla")
        totales = self._cell_with_value(worksheet, "Totales")
        self.assertIsNotNone(datos)
        self.assertIsNotNone(totales)
        self.assertLess(datos.row, totales.row)
        self.assertIsNone(self._cell_with_value(worksheet, "Resumen"))
        recon = payload["recon"]
        for label, key in (
            ("Efectivo a entregar", "cash_due"),
            ("Cobrado por otros medios", "other_collected"),
            ("Comisión a Pagar al Vendedor", "commission_owed"),
        ):
            with self.subTest(label=label):
                header = self._cell_with_value(worksheet, label)
                self.assertIsNotNone(header)
                self.assertEqual(header.row, 7)
                value = worksheet.cell(header.row + 1, header.column)
                self.assertTrue(str(value.value).startswith("="), value.value)
                self.assertAlmostEqual(
                    float(cached.cell(value.row, value.column).value),
                    recon[key],
                    places=2,
                )
        self.assertLess(8, datos.row)
        self.assertLess(totales.row, self._cell_with_value(worksheet, "Clientes").row)

    def test_xlsx_mismatch_names_both_amounts_and_gap(self):
        """Mismatch MUST print the allocated total, the voucher total and the gap."""
        sheet, partner = self._populated_sheet(amount=100.0)
        self._create_voucher(partner, sheet, amount=40.0)
        content, payload = self._render_xlsx(sheet)
        self.assertTrue(payload["mismatch"])
        worksheet = self._load(content).active
        warning = next(
            (
                cell
                for row in worksheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and "no coincide" in cell.value
            ),
            None,
        )
        self.assertIsNotNone(warning)
        self.assertEqual(warning.row, 9)
        for amount in ("100.00", "140.00", "40.00"):
            self.assertIn(amount, warning.value)
        clean_seller = create_internal_seller(self.env, "xlsx_no_mismatch")
        ensure_enabled_seller(self.env, clean_seller)
        clean, _partner = self._populated_sheet(
            seller=clean_seller, partner_name="Sin Diferencia SA"
        )
        clean_content, _payload = self._render_xlsx(clean)
        clean_ws = self._load(clean_content).active
        self.assertFalse(
            any(
                isinstance(cell.value, str) and "no coincide" in cell.value
                for row in clean_ws.iter_rows()
                for cell in row
            )
        )

    def test_xlsx_tendered_deposits_are_explained(self):
        """Customer and seller deposits MUST be told apart under the composition."""
        sheet, _partner = self._populated_sheet()
        content, _payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        header = self._cell_with_value(worksheet, "Depósitos de clientes")
        self.assertIsNotNone(header)
        self.assertIsNotNone(self._cell_with_value(worksheet, "Depósitos del vendedor"))
        self.assertIsNone(self._cell_with_value(worksheet, "Depósitos"))
        note = worksheet.cell(header.row + 2, 1).value or ""
        self.assertIn("el cliente depositó en la cuenta de la empresa", note)
        self.assertIn("el vendedor depositó en el banco el efectivo cobrado", note)

    def test_xlsx_clientes_has_no_motivo_column(self):
        """Clientes MUST end at Deuda sin cobrar: no Motivo del Remanente column."""
        sheet, _partner = self._populated_sheet()
        content, _payload = self._render_xlsx(sheet)
        worksheet = self._load(content).active
        self.assertIsNone(self._cell_with_value(worksheet, "Motivo del Remanente"))
        saldo = self._cell_with_value(worksheet, "Saldo Total")
        header = [worksheet.cell(saldo.row, col).value for col in range(1, 9)]
        self.assertEqual(
            header,
            [
                "Código",
                "Cliente",
                "Saldo Total",
                "Cobrado",
                "Aplicado",
                "Deuda sin cobrar",
                None,
                None,
            ],
        )

    def test_xlsx_has_only_the_planilla_sheet(self):
        """The Documentos and Comprobantes flat sheets MUST be gone."""
        sheet, _partner = self._populated_sheet()
        content, _payload = self._render_xlsx(sheet)
        workbook = self._load(content)
        self.assertEqual(workbook.sheetnames, ["Planilla de Cobranza"])
        self.assertEqual(workbook.active.title, "Planilla de Cobranza")
        self.assertEqual(workbook.active.column_dimensions["B"].width, 32)
